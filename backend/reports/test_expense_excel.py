"""WF3 single-sheet Excel export tests."""
from datetime import datetime
from io import BytesIO

from backend.testing import ASGITestClient as TestClient
from openpyxl import load_workbook

from backend.backends.records import RecordStore
from backend.main import app
from backend.reports.expense_excel import HEADERS, export_expense_workbook


def _claim(store: RecordStore, *, status: str, **overrides):
    data = {
        "employee_name": "Alice Tan",
        "vendor": "Café Aurora",
        "date": "2026-08-01",
        "amount": 75.48,
        "currency": "EUR",
        "amount_gbp": 64.23,
        "category": "meals",
        "business_purpose": "Client lunch",
        "payment_method": "corporate_card",
        "submitted_at": "2026-08-02T09:15:00+00:00",
    }
    data.update(overrides)
    return store.create("submissions", "expense_claim", data, status=status)


def test_workbook_is_one_plain_sheet_and_excludes_pre_submission_drafts():
    store = RecordStore(":memory:")
    _claim(store, status="pending_extraction", vendor="Draft vendor")
    approved = _claim(
        store,
        status="approved",
        department=None,
        reviewed_by="chen",
        decided_at="2026-08-03T14:30:00Z",
    )

    workbook = load_workbook(BytesIO(export_expense_workbook(store.list("submissions"))))
    assert workbook.sheetnames == ["Receipts"]
    sheet = workbook["Receipts"]
    assert [cell.value for cell in sheet[1]] == list(HEADERS)
    assert sheet.max_row == 2
    assert sheet["A2"].value == approved.id
    assert sheet["B2"].value == "Alice Tan"
    assert sheet["C2"].value == "Product"  # derived for legacy records
    assert sheet["D2"].value == "Café Aurora"
    assert sheet["E2"].value == datetime(2026, 8, 1)
    assert sheet["I2"].value == 75.48
    assert sheet["J2"].value == "EUR"
    assert sheet["K2"].value == 64.23
    assert sheet["N2"].value == "Chen Wei"
    assert sheet["O2"].value == datetime(2026, 8, 3, 14, 30)
    assert sheet.freeze_panes is None
    assert sheet.sheet_view.showGridLines is None
    assert list(sheet.tables) == []
    assert sheet.auto_filter.ref is None
    assert sheet["A1"].fill.fill_type is None
    assert sheet["A1"].font.bold is False


def test_rejected_claim_exports_reason_and_legacy_timestamps():
    store = RecordStore(":memory:")
    rejected = _claim(
        store,
        status="rejected",
        submitted_at=None,
        reviewed_by="fiona",
        decision_reason="Personal purchase",
    )
    workbook = load_workbook(BytesIO(export_expense_workbook([rejected])))
    sheet = workbook["Receipts"]
    assert sheet["L2"].value == datetime.fromisoformat(rejected.created_at).replace(tzinfo=None)
    assert sheet["N2"].value == "Fiona Reyes"
    assert sheet["O2"].value == datetime.fromisoformat(rejected.updated_at).replace(tzinfo=None)
    assert sheet["P2"].value == "Personal purchase"


def test_user_text_is_not_exported_as_an_excel_formula():
    store = RecordStore(":memory:")
    claim = _claim(store, status="submitted", vendor="=HYPERLINK(\"bad\")")
    workbook = load_workbook(BytesIO(export_expense_workbook([claim])), data_only=False)
    vendor = workbook["Receipts"]["D2"]
    assert vendor.data_type == "s"
    assert vendor.value == "'=HYPERLINK(\"bad\")"
    assert not list(workbook["Receipts"].iter_rows(values_only=False))[1][3].data_type == "f"


def test_export_endpoint_returns_downloadable_xlsx():
    response = TestClient(app).get("/api/expense/ledger/export.xlsx")
    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "wf3_receipts_" in response.headers["content-disposition"]
    assert response.content.startswith(b"PK")
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    assert workbook.sheetnames == ["Receipts"]
